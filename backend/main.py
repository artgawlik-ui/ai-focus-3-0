import io, os, json, secrets, uuid, logging, hashlib, unicodedata
from datetime import datetime, timezone, timedelta, date
from pathlib import Path
from urllib.parse import urlencode
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File, Query
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy import select, func, or_, update
from sqlalchemy.exc import IntegrityError
from argon2.exceptions import VerifyMismatchError, InvalidHashError
from openpyxl import Workbook, load_workbook
from backend.db import Session, User, LoginSession, LoginAttempt, Salon, Activity, Appointment, Job, Training, now
from backend.security import database, current, manager, digest, expires, secure, ph, public_user
from backend.schemas import *

logging.basicConfig(level=logging.INFO)
log = logging.getLogger('focus')
app = FastAPI(title='Ai Focus 3.0', version='3.0.0')
uid = lambda: str(uuid.uuid4())

@app.exception_handler(IntegrityError)
async def integrity_error(request, exc):
    return JSONResponse(status_code=409, content={'detail':'Duplikat danych lub konflikt zapisu'})


def serial(obj): return {c.name:getattr(obj,c.name) for c in obj.__table__.columns}
def key(name, city):
    text = unicodedata.normalize('NFKD', name+'|'+city).casefold()
    return digest(''.join(c for c in text if c.isalnum() or c=='|'))
def scope(stmt,user): return stmt if user.role=='manager' else stmt.where(Salon.owner_id==user.id)
def salon_access(db,id,user):
    salon=db.scalar(scope(select(Salon).where(Salon.id==id),user))
    if not salon: raise HTTPException(404,'Salon nie istnieje')
    return salon
def record(db,salon,user,kind,body):
    db.flush()
    db.add(Activity(id=uid(),salon_id=salon.id,user_id=user.id,kind=kind,body=body))

def commit(db):
    try: db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409,'Duplikat danych lub konflikt zapisu')

@app.middleware('http')
async def headers(request: Request, call_next):
    request_id=secrets.token_hex(8)
    if request.method not in ('GET','HEAD','OPTIONS'):
        origin=request.headers.get('origin')
        allowed=os.getenv('APP_ORIGIN',str(request.base_url).rstrip('/'))
        if origin and origin != allowed:
            return Response('Niedozwolone źródło',status_code=403)
    response=await call_next(request)
    response.headers.update({'X-Request-ID':request_id,'X-Content-Type-Options':'nosniff','X-Frame-Options':'DENY','Referrer-Policy':'strict-origin-when-cross-origin','Permissions-Policy':'camera=(), microphone=(), geolocation=()','Content-Security-Policy':"default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'self'; frame-ancestors 'none'; form-action 'self'"})
    if request.url.path.startswith('/api'): response.headers['Cache-Control']='no-store'
    if secure: response.headers['Strict-Transport-Security']='max-age=31536000; includeSubDomains'
    log.info('request id=%s method=%s status=%s',request_id,request.method,response.status_code)
    return response

@app.get('/api/health')
def health(db=Depends(database)):
    db.execute(select(func.count()).select_from(User))
    return {'status':'ok','version':'3.0.0'}

@app.post('/api/auth/login')
def login(payload:LoginIn, request:Request, response:Response, db=Depends(database)):
    email=payload.email.strip().lower()
    # Email-based durable throttling works across processes, without trusting spoofable proxy headers.
    bucket=digest(email)
    cutoff=(datetime.now(timezone.utc)-timedelta(minutes=15)).isoformat()
    attempts=db.scalar(select(func.count()).select_from(LoginAttempt).where(LoginAttempt.key==bucket,LoginAttempt.at>cutoff))
    if attempts>=8: raise HTTPException(429,'Zbyt wiele prób. Spróbuj za 15 minut.')
    user=db.scalar(select(User).where(User.email==email))
    valid=False
    try: valid=ph.verify(user.password_hash if user else DUMMY_HASH,payload.password)
    except (VerifyMismatchError,InvalidHashError): pass
    if not valid or not user or not user.active:
        db.add(LoginAttempt(id=uid(),key=bucket,at=now())); db.commit()
        raise HTTPException(401,'Nieprawidłowy e-mail lub hasło')
    token=secrets.token_urlsafe(48); csrf=secrets.token_urlsafe(32)
    db.add(LoginSession(token_hash=digest(token),user_id=user.id,csrf=csrf,expires=expires()))
    db.commit()
    response.set_cookie('focus_session',token,httponly=True,secure=secure,samesite='strict',max_age=43200,path='/')
    return {'user':public_user(user),'csrf':csrf}
DUMMY_HASH=ph.hash(secrets.token_urlsafe(32))

@app.get('/api/auth/me')
def me(request:Request,user=Depends(current),db=Depends(database)):
    session=db.get(LoginSession,digest(request.cookies['focus_session']))
    return {'user':public_user(user),'csrf':session.csrf}
@app.post('/api/auth/logout')
def logout(request:Request,response:Response,user=Depends(current),db=Depends(database)):
    db.delete(db.get(LoginSession,digest(request.cookies['focus_session']))); db.commit()
    response.delete_cookie('focus_session',path='/',secure=secure,httponly=True,samesite='strict')
    return {'ok':True}
@app.get('/api/users')
def users(user=Depends(manager),db=Depends(database)):
    return [public_user(u) for u in db.scalars(select(User).order_by(User.name))]
@app.post('/api/users',status_code=201)
def add_user(payload:UserIn,user=Depends(manager),db=Depends(database)):
    if db.scalar(select(func.count()).select_from(User).where(User.active==1))>=6:
        raise HTTPException(409,'Limit sześciu aktywnych kont')
    new=User(id=uid(),email=payload.email.strip().lower(),name=payload.name,password_hash=ph.hash(payload.password),role=payload.role)
    db.add(new); commit(db)
    return public_user(new)
@app.get('/api/salons')
def salons(q:str='',priority:str='',city:str='',stage:str='',limit:int=Query(100,ge=1,le=500),offset:int=Query(0,ge=0),user=Depends(current),db=Depends(database)):
    stmt=scope(select(Salon),user)
    if q: stmt=stmt.where(or_(Salon.name.ilike('%'+q+'%'),Salon.city.ilike('%'+q+'%'),Salon.phone.ilike('%'+q+'%')))
    if priority: stmt=stmt.where(Salon.priority==priority)
    if city: stmt=stmt.where(Salon.city==city)
    if stage: stmt=stmt.where(Salon.stage==stage)
    total=db.scalar(select(func.count()).select_from(stmt.subquery()))
    return {'items':[serial(s) for s in db.scalars(stmt.order_by(Salon.updated_at.desc()).limit(limit).offset(offset))],'total':total}
@app.post('/api/salons',status_code=201)
def add_salon(payload:SalonIn,user=Depends(current),db=Depends(database)):
    salon=Salon(id=uid(),owner_id=user.id,**payload.model_dump(),dedup_key=key(payload.name,payload.city))
    db.add(salon); record(db,salon,user,'created','Dodano salon'); commit(db)
    return serial(salon)
@app.get('/api/salons/{id}')
def salon_detail(id:str,user=Depends(current),db=Depends(database)):
    return serial(salon_access(db,id,user))
@app.put('/api/salons/{id}')
def edit_salon(id:str,payload:SalonUpdate,user=Depends(current),db=Depends(database)):
    salon=salon_access(db,id,user)
    values=payload.model_dump(exclude={'version'})
    changes=[f'{k}: {getattr(salon,k)} → {v}' for k,v in values.items() if getattr(salon,k)!=v]
    result=db.execute(update(Salon).where(Salon.id==id,Salon.version==payload.version).values(**values,version=payload.version+1,updated_at=now(),dedup_key=key(payload.name,payload.city)))
    if result.rowcount!=1: db.rollback(); raise HTTPException(409,'Rekord zmieniono w innym oknie. Odśwież dane.')
    record(db,salon,user,'updated','; '.join(changes) or 'Zapis bez zmian'); commit(db)
    db.refresh(salon); return serial(salon)
@app.get('/api/salons/{id}/activities')
def activities(id:str,user=Depends(current),db=Depends(database)):
    salon_access(db,id,user)
    return [serial(a) for a in db.scalars(select(Activity).where(Activity.salon_id==id).order_by(Activity.created_at.desc()))]
@app.post('/api/salons/{id}/activities',status_code=201)
def add_activity(id:str,payload:ActivityIn,user=Depends(current),db=Depends(database)):
    salon=salon_access(db,id,user); record(db,salon,user,payload.kind,payload.body); db.commit()
    return {'ok':True}

@app.get('/api/export.xlsx')
def export(user=Depends(current),db=Depends(database)):
    wb=Workbook(); ws=wb.active; ws.title='Salony'
    fields=list(SalonIn.model_fields); ws.append(fields)
    for s in db.scalars(scope(select(Salon),user)):
        row=[]
        for field in fields:
            value=getattr(s,field)
            if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')): value="'"+value
            row.append(value)
        ws.append(row)
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="focus3-salony.xlsx"'})

@app.post('/api/import')
async def import_excel(file:UploadFile=File(...),dry_run:bool=True,user=Depends(current),db=Depends(database)):
    import zipfile
    raw=await file.read(5*1024*1024+1)
    if len(raw)>5*1024*1024: raise HTTPException(413,'Limit pliku: 5 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>30*1024*1024: raise ValueError('Archiwum zbyt duże')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        rows=wb.active.iter_rows(values_only=True); fields=next(rows)
        if not fields or len(set(fields))!=len(fields) or not {'name','city'}.issubset(fields): raise ValueError('Wymagane unikalne nagłówki name i city')
        if any(f not in SalonIn.model_fields for f in fields): raise ValueError('Nieznane kolumny; użyj wzoru eksportu')
        accepted=[]; errors=[]; seen=set(db.scalars(select(Salon.dedup_key)))
        for idx,row in enumerate(rows,2):
            if idx>2001: raise ValueError('Limit 2000 wierszy')
            if all(v is None for v in row): continue
            try:
                data={f:v for f,v in zip(fields,row) if v is not None}
                if any(isinstance(v,str) and v.startswith('=') for v in data.values()): raise ValueError('Formuły nie są dozwolone')
                payload=SalonIn.model_validate(data); dedup=key(payload.name,payload.city)
                if dedup in seen: raise ValueError('Duplikat nazwy i miasta')
                seen.add(dedup); accepted.append(payload)
            except Exception: errors.append({'row':idx,'error':'Nieprawidłowe dane lub duplikat'})
        wb.close()
    except Exception as e: raise HTTPException(422,'Nieprawidłowy plik XLSX: '+str(e)[:160])
    if not dry_run:
        if errors: raise HTTPException(422,{'message':'Import anulowany w całości; popraw błędy','errors':errors})
        for p in accepted:
            s=Salon(id=uid(),owner_id=user.id,**p.model_dump(),dedup_key=key(p.name,p.city)); db.add(s); record(db,s,user,'imported','Import Excel')
        commit(db)
    return {'dry_run':dry_run,'accepted':len(accepted),'errors':errors,'saved':0 if dry_run else len(accepted)}

@app.get('/api/appointments')
def appointments(day:str='',user=Depends(current),db=Depends(database)):
    stmt=select(Appointment,Salon).join(Salon,Salon.id==Appointment.salon_id)
    if user.role!='manager': stmt=stmt.where(Appointment.user_id==user.id)
    if day:
        try: date.fromisoformat(day)
        except ValueError: raise HTTPException(422,'Nieprawidłowy dzień')
        # Europe/Warsaw day boundaries, including DST.
        from zoneinfo import ZoneInfo
        start=datetime.combine(date.fromisoformat(day),datetime.min.time(),ZoneInfo('Europe/Warsaw'))
        stmt=stmt.where(Appointment.starts_at>=start.astimezone(timezone.utc).isoformat(),Appointment.starts_at<(start+timedelta(days=1)).astimezone(timezone.utc).isoformat())
    return [{**serial(a),'salon_name':s.name,'city':s.city,'address':s.address} for a,s in db.execute(stmt.order_by(Appointment.starts_at))]
@app.post('/api/appointments',status_code=201)
def add_appointment(payload:AppointmentIn,user=Depends(current),db=Depends(database)):
    salon=salon_access(db,payload.salon_id,user)
    if payload.starts_at.tzinfo is None: raise HTTPException(422,'Termin wymaga strefy czasowej')
    start=payload.starts_at.astimezone(timezone.utc); end=start+timedelta(minutes=payload.minutes)
    owner=salon.owner_id
    db.scalar(select(User).where(User.id==owner).with_for_update())
    overlap=db.scalar(select(Appointment).where(Appointment.user_id==owner,Appointment.status.in_(['proposed','confirmed']),Appointment.starts_at<end.isoformat(),Appointment.ends_at>start.isoformat()))
    if overlap: raise HTTPException(409,'Termin koliduje z inną wizytą')
    a=Appointment(id=uid(),salon_id=salon.id,user_id=owner,starts_at=start.isoformat(),ends_at=end.isoformat(),status=payload.status,notes=payload.notes)
    db.add(a); record(db,salon,user,'appointment','Zaplanowano wizytę '+a.starts_at+' ('+a.status+')'); db.commit(); return serial(a)
@app.patch('/api/appointments/{id}')
def appointment_status(id:str,payload:AppointmentStatus,user=Depends(current),db=Depends(database)):
    a=db.get(Appointment,id)
    if not a or (user.role!='manager' and a.user_id!=user.id): raise HTTPException(404,'Brak wizyty')
    db.scalar(select(User).where(User.id==a.user_id).with_for_update())
    if payload.status in ('confirmed','proposed'):
        overlap=db.scalar(select(Appointment).where(Appointment.id!=id,Appointment.user_id==a.user_id,Appointment.status.in_(['proposed','confirmed']),Appointment.starts_at<a.ends_at,Appointment.ends_at>a.starts_at))
        if overlap: raise HTTPException(409,'Termin koliduje z inną wizytą')
    a.status=payload.status; record(db,db.get(Salon,a.salon_id),user,'visit' if a.status=='completed' else 'appointment','Status wizyty: '+a.status); db.commit(); return serial(a)
@app.get('/api/routes')
def routes(day:str,user=Depends(current),db=Depends(database)):
    visits=[v for v in appointments(day,user,db) if v['status']=='confirmed' and v['address']]
    stops=[v['address']+', '+v['city']+', Polska' for v in visits]
    params={'api':'1','travelmode':'driving'}
    if stops:
        params['destination']=stops[-1]
        if len(stops)>1: params['origin']=stops[0]
        if len(stops)>2: params['waypoints']='|'.join(stops[1:-1])
    return {'visits':visits,'url':'https://www.google.com/maps/dir/?'+urlencode(params) if stops else None,'note':'Kolejność według godzin wizyt. Brak optymalizacji ruchu drogowego.'}

@app.get('/api/dashboard')
def dashboard(days:int=Query(30,ge=1,le=366),user=Depends(current),db=Depends(database)):
    salons=list(db.scalars(scope(select(Salon),user))); ids=[s.id for s in salons]
    cutoff=(datetime.now(timezone.utc)-timedelta(days=days)).isoformat()
    activities=list(db.scalars(select(Activity).where(Activity.salon_id.in_(ids),Activity.created_at>=cutoff)))
    pipeline={stage:sum(s.stage==stage for s in salons) for stage in ['new','contacted','discovery','offer','negotiation','won','lost']}
    weights={'new':.05,'contacted':.1,'discovery':.25,'offer':.5,'negotiation':.75,'won':1,'lost':0}
    due=[serial(s) for s in salons if s.next_contact and s.next_contact<=date.today().isoformat() and s.status!='do_not_contact']
    return {'total':len(salons),'active':sum(s.status=='active' for s in salons),'potential':sum(s.potential for s in salons),'weighted_forecast':round(sum(s.potential*weights[s.stage] for s in salons)),'forecast_note':'Model orientacyjny: stałe wagi etapów, bez danych zamówień. Potencjał miesięczny w PLN.','pipeline':pipeline,'activities':len(activities),'days':days,'conversion':round(100*pipeline['won']/max(1,pipeline['won']+pipeline['lost']),1),'follow_ups':due,'team':[{'name':u.name,'salons':sum(s.owner_id==u.id for s in salons),'activities':sum(a.user_id==u.id for a in activities)} for u in db.scalars(select(User).where(User.active==1))] if user.role=='manager' else []}
@app.get('/api/integrations')
def integrations(user=Depends(current)):
    return {'instagram':{'connected':False,'outbound':False},'google_calendar':{'connected':False,'mode':'manual_link'},'livespace':{'connected':False,'mode':'adapter_schema_ready'},'agent':{'mode':'openai' if os.getenv('OPENAI_API_KEY') else 'rules','outbound':False},'discovery':{'provider':'OpenStreetMap','connected':True,'note':'Publiczne wpisy wymagają weryfikacji; nie stanowią zgody na kontakt.'}}
@app.post('/api/agent/jobs',status_code=202)
def queue_job(payload:JobIn,user=Depends(current),db=Depends(database)):
    salon=salon_access(db,payload.salon_id,user)
    if salon.status=='do_not_contact': raise HTTPException(409,'Salon ma zablokowany kontakt')
    db.scalar(select(Salon).where(Salon.id==salon.id).with_for_update())
    pending=db.scalar(select(Job).where(Job.salon_id==salon.id,Job.status.in_(['queued','running'])))
    if pending: raise HTTPException(409,'Analiza salonu już oczekuje')
    job=Job(id=uid(),salon_id=salon.id,user_id=user.id,evidence=payload.evidence); db.add(job); db.commit(); return serial(job)
@app.get('/api/agent/jobs')
def jobs(user=Depends(current),db=Depends(database)):
    stmt=select(Job)
    if user.role!='manager': stmt=stmt.where(Job.user_id==user.id)
    return [serial(j) for j in db.scalars(stmt.order_by(Job.created_at.desc()).limit(100))]
@app.get('/api/discovery')
def discovery(city:str=Query(min_length=2,max_length=80),user=Depends(current)):
    import httpx
    # JSON quoting prevents Overpass query injection. No arbitrary URL fetching.
    city_json=json.dumps(city)
    query=f'[out:json][timeout:20];area["name"={city_json}]["boundary"="administrative"]["admin_level"~"6|7|8"]->.a;nwr["shop"="hairdresser"](area.a);out center tags 40;'
    try:
        result=httpx.post('https://overpass-api.de/api/interpreter',data={'data':query},timeout=30,headers={'User-Agent':'AiFocus3/3.0 CRM public-business-research'})
        result.raise_for_status(); elements=result.json().get('elements',[])
    except Exception: raise HTTPException(503,'Wyszukiwarka OSM chwilowo niedostępna; spróbuj później')
    return {'source':'© OpenStreetMap contributors, ODbL','items':[{'name':e.get('tags',{}).get('name','Salon bez nazwy'),'city':city,'address':' '.join(filter(None,[e.get('tags',{}).get('addr:street'),e.get('tags',{}).get('addr:housenumber')])),'phone':e.get('tags',{}).get('phone',''),'source_url':f'https://www.openstreetmap.org/{e["type"]}/{e["id"]}'} for e in elements]}

@app.post('/api/auth/password')
def password_change(payload:PasswordIn,request:Request,user=Depends(current),db=Depends(database)):
    try: ph.verify(user.password_hash,payload.current_password)
    except (VerifyMismatchError,InvalidHashError): raise HTTPException(403,'Nieprawidłowe obecne hasło')
    user.password_hash=ph.hash(payload.new_password)
    from sqlalchemy import delete
    active_hash=digest(request.cookies['focus_session'])
    db.execute(delete(LoginSession).where(LoginSession.user_id==user.id,LoginSession.token_hash!=active_hash))
    db.commit(); return {'ok':True}

@app.post('/api/salons/{id}/assign')
def assign_salon(id:str,payload:AssignmentIn,user=Depends(manager),db=Depends(database)):
    salon=salon_access(db,id,user); owner=db.get(User,payload.owner_id)
    if not owner or not owner.active: raise HTTPException(422,'Nieprawidłowy opiekun')
    salon.owner_id=owner.id; salon.version+=1; salon.updated_at=now()
    record(db,salon,user,'assignment','Przypisano opiekuna: '+owner.name); db.commit()
    return serial(salon)

@app.get('/api/trainings')
def trainings(user=Depends(current),db=Depends(database)):
    return [serial(t) for t in db.scalars(select(Training).order_by(Training.starts_at))]
@app.post('/api/trainings',status_code=201)
def add_training(payload:TrainingIn,user=Depends(manager),db=Depends(database)):
    if payload.starts_at.tzinfo is None: raise HTTPException(422,'Termin wymaga strefy czasowej')
    t=Training(id=uid(),title=payload.title,starts_at=payload.starts_at.astimezone(timezone.utc).isoformat(),minutes=payload.minutes,notes=payload.notes,created_by=user.id)
    db.add(t); db.commit(); return serial(t)

@app.post('/api/users/{id}/deactivate')
def deactivate_user(id:str,user=Depends(manager),db=Depends(database)):
    target=db.get(User,id)
    if not target: raise HTTPException(404,'Brak użytkownika')
    if target.role=='manager' and target.active:
        active_managers=db.scalar(select(func.count()).select_from(User).where(User.role=='manager',User.active==1))
        if active_managers<=1: raise HTTPException(409,'Nie można wyłączyć ostatniego menedżera')
    target.active=0
    from sqlalchemy import delete
    db.execute(delete(LoginSession).where(LoginSession.user_id==id));db.commit()
    return {'ok':True}

@app.post('/api/admin/backup')
def download_backup(payload:BackupIn,user=Depends(manager)):
    import tempfile, shutil
    from starlette.background import BackgroundTask
    from scripts.backup import backup
    if not os.getenv('DATABASE_URL','').startswith(('postgres://','postgresql://','postgresql+psycopg://')):
        raise HTTPException(409,'Kopie bazy wymagają PostgreSQL')
    directory=tempfile.mkdtemp(prefix='focus3-backup-'); path=Path(directory)/'focus3.enc'
    try: backup(str(path),phrase=payload.passphrase)
    except Exception:
        shutil.rmtree(directory,ignore_errors=True)
        raise HTTPException(503,'Kopia nie powiodła się. Sprawdź klienta PostgreSQL i dostępność bazy.')
    return FileResponse(path,filename='focus3-backup.enc',media_type='application/octet-stream',background=BackgroundTask(shutil.rmtree,directory,ignore_errors=True),headers={'Cache-Control':'no-store'})

DIST=Path(__file__).resolve().parent.parent/'frontend'/'dist'
if DIST.exists():
    app.mount('/assets',StaticFiles(directory=DIST/'assets'),name='assets')
    @app.get('/{path:path}',include_in_schema=False)
    def spa(path:str):
        if path.startswith('api/'): raise HTTPException(404)
        file=(DIST/path).resolve()
        if file != DIST.resolve() and DIST.resolve() not in file.parents: raise HTTPException(404)
        return FileResponse(file if file.is_file() else DIST/'index.html',headers={'Cache-Control':'no-cache'})
