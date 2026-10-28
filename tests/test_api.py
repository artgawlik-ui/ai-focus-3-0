import os, uuid, io
os.environ.setdefault('DATABASE_URL','sqlite:////tmp/focus3-tests.db')
os.environ['APP_ENV']='test'
from sqlalchemy.engine import make_url
test_url=make_url(os.environ['DATABASE_URL'])
if not (test_url.database=='focus3_test' and test_url.host in ('localhost','127.0.0.1','postgres')) and str(test_url)!='sqlite:////tmp/focus3-tests.db':
    raise RuntimeError('Refusing destructive tests outside isolated focus3_test database')
from fastapi.testclient import TestClient
from sqlalchemy import select
from backend.db import Base, engine, Session, User, Job, Activity
from backend.security import ph
from backend.main import app
from backend.worker import run_once
import pytest

@pytest.fixture(autouse=True)
def clean():
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine)
    with Session() as db:
        for id,role in [('manager','manager'),('rep1','sales'),('rep2','sales')]:
            db.add(User(id=id,name=id,email=id+'@focus.test',password_hash=ph.hash('TestPassword123!'),role=role))
        db.commit()

def client(id='manager'):
    c=TestClient(app); r=c.post('/api/auth/login',json={'email':id+'@focus.test','password':'TestPassword123!'})
    assert r.status_code==200,r.text
    c.headers['X-CSRF-Token']=r.json()['csrf'];return c

def add(c,name='Salon Testowy',**kwargs):
    r=c.post('/api/salons',json={'name':name,'city':'Głogów',**kwargs}); assert r.status_code==201,r.text;return r.json()

def test_auth_csrf_origin_and_password():
    c=TestClient(app)
    assert c.get('/api/salons').status_code==401
    assert c.post('/api/auth/login',json={'email':'manager@focus.test','password':'bad'}).status_code==401
    c=client();assert c.post('/api/salons',json={'name':'Salon','city':'Głogów'},headers={'X-CSRF-Token':''}).status_code==403
    assert c.post('/api/salons',json={},headers={'Origin':'https://evil.test'}).status_code==403
    assert c.post('/api/auth/password',json={'current_password':'TestPassword123!','new_password':'NewPassword123!!'}).status_code==200
    assert c.post('/api/auth/logout').status_code==200
    assert c.get('/api/auth/me').status_code==401
    assert c.post('/api/auth/login',json={'email':'manager@focus.test','password':'NewPassword123!!'}).status_code==200

def test_roles_ownership_reassignment():
    a=client('rep1'); b=client('rep2'); m=client(); s=add(a)
    assert b.get('/api/salons').json()['total']==0
    assert b.get('/api/salons/'+s['id']).status_code==404
    assert b.post('/api/agent/jobs',json={'salon_id':s['id']}).status_code==404
    assert a.get('/api/users').status_code==403
    assert m.get('/api/salons').json()['total']==1
    assert m.post('/api/salons/'+s['id']+'/assign',json={'owner_id':'rep2'}).status_code==200
    assert a.get('/api/salons/'+s['id']).status_code==404
    assert b.get('/api/salons/'+s['id']).status_code==200

def test_duplicate_optimistic_lock_and_audit():
    c=client(); s=add(c)
    assert c.post('/api/salons',json={'name':'salon testowy','city':'Głogów'}).status_code==409
    payload={k:v for k,v in s.items() if k in __import__('backend.schemas',fromlist=['SalonUpdate']).SalonUpdate.model_fields};payload['stage']='offer'
    assert c.put('/api/salons/'+s['id'],json=payload).status_code==200
    assert c.put('/api/salons/'+s['id'],json=payload).status_code==409
    assert len(c.get('/api/salons/'+s['id']+'/activities').json())==2

def test_import_atomic_dry_run_export_sanitization():
    from openpyxl import Workbook,load_workbook
    c=client();add(c,name='=INJECTION')
    workbook=load_workbook(io.BytesIO(c.get('/api/export.xlsx').content)); assert workbook.active['A2'].value.startswith("'")
    wb=Workbook();ws=wb.active;ws.append(['name','city']);ws.append(['Nowy Salon','Lubin']);ws.append(['Nowy Salon','Lubin']);raw=io.BytesIO();wb.save(raw)
    r=c.post('/api/import?dry_run=true',files={'file':('test.xlsx',raw.getvalue())});assert r.json()['accepted']==1;assert len(r.json()['errors'])==1
    assert c.post('/api/import?dry_run=false',files={'file':('test.xlsx',raw.getvalue())}).status_code==422
    assert c.get('/api/salons').json()['total']==1
    ws.delete_rows(3);raw=io.BytesIO();wb.save(raw)
    assert c.post('/api/import?dry_run=false',files={'file':('test.xlsx',raw.getvalue())}).json()['saved']==1
    assert c.get('/api/salons').json()['total']==2

def test_appointments_collisions_routes_status():
    c=client();s=add(c,address='Rynek 1')
    p={'salon_id':s['id'],'starts_at':'2027-01-20T10:00:00+01:00','minutes':45,'status':'confirmed'}
    r=c.post('/api/appointments',json=p);assert r.status_code==201
    assert c.post('/api/appointments',json=p).status_code==409
    assert c.get('/api/routes?day=2027-01-20').json()['url'].startswith('https://www.google.com/maps/dir/?')
    assert c.patch('/api/appointments/'+r.json()['id'],json={'status':'completed'}).status_code==200
    assert c.get('/api/routes?day=2027-01-20').json()['url'] is None

def test_agent_durable_job_history_and_no_send():
    c=client();s=add(c)
    c.post('/api/salons/'+s['id']+'/activities',json={'kind':'call','body':'Rozmowa dzisiaj'})
    r=c.post('/api/agent/jobs',json={'salon_id':s['id']});assert r.status_code==202
    assert c.post('/api/agent/jobs',json={'salon_id':s['id']}).status_code==409
    assert run_once()
    import json
    result=json.loads(c.get('/api/agent/jobs').json()[0]['result'])
    assert result['sent'] is False
    assert 'Odczekaj' in result['next_step']
    assert not run_once()
    assert any(a['kind']=='agent_draft' for a in c.get('/api/salons/'+s['id']+'/activities').json())

def test_contact_optout_and_dashboard_scope():
    c=client('rep1');s=add(c,status='do_not_contact',next_contact='2020-01-01',potential=2000)
    assert c.post('/api/agent/jobs',json={'salon_id':s['id']}).status_code==409
    assert c.get('/api/dashboard').json()['follow_ups']==[]
    assert client('rep2').get('/api/dashboard').json()['total']==0
    assert c.get('/api/dashboard').json()['weighted_forecast']==100

def test_login_rate_limit():
    c=TestClient(app)
    for _ in range(8):assert c.post('/api/auth/login',json={'email':'x@focus.test','password':'bad'}).status_code==401
    assert c.post('/api/auth/login',json={'email':'x@focus.test','password':'bad'}).status_code==429

def test_input_validation_security_headers_and_persistence():
    c=client()
    assert c.post('/api/salons',json={'name':'A','city':'','potential':-1}).status_code==422
    s=add(c)
    assert client().get('/api/salons/'+s['id']).json()['name']==s['name']
    assert "frame-ancestors 'none'" in c.get('/api/health').headers['content-security-policy']
    assert c.get('/api/salons').headers['cache-control']=='no-store'

def test_training_manager_permissions():
    m=client();r=client('rep1')
    payload={'title':'Szkolenie NAK Hair','starts_at':'2027-02-01T10:00:00+01:00','minutes':60,'notes':'Portfolio i konsultacje'}
    assert r.post('/api/trainings',json=payload).status_code==403
    assert m.post('/api/trainings',json=payload).status_code==201
    assert r.get('/api/trainings').json()[0]['title']==payload['title']

def test_deactivate_sessions_and_protect_last_manager():
    m=client(); r=client('rep1')
    assert r.post('/api/users/rep2/deactivate').status_code==403
    assert m.post('/api/users/manager/deactivate').status_code==409
    assert m.post('/api/users/rep1/deactivate').status_code==200
    assert r.get('/api/auth/me').status_code==401
    assert r.post('/api/auth/login',json={'email':'rep1@focus.test','password':'TestPassword123!'}).status_code==401

def test_backup_permission_and_passphrase_validation():
    assert client('rep1').post('/api/admin/backup',json={'passphrase':'x'*30}).status_code==403
    assert client().post('/api/admin/backup',json={'passphrase':'short'}).status_code==422

def test_integrity_is_manager_only():
    assert TestClient(app).get('/api/admin/integrity').status_code==401
    assert client('rep1').get('/api/admin/integrity').status_code==403

def test_body_limits_before_parser():
    c=TestClient(app)
    assert c.post('/api/auth/login',content=b'x'*(256*1024+1),headers={'Content-Type':'application/json'}).status_code==413
    assert c.post('/api/import',content=b'x',headers={'Content-Length':str(7*1024*1024)}).status_code==413
    assert c.post('/api/auth/login',content=iter([b'x'*150000,b'x'*150000]),headers={'Content-Type':'application/json'}).status_code==413
