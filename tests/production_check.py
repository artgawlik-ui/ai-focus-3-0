"""Authorized production checks against synthetic records only; credentials stay outside git."""
import io, json, os, time
from pathlib import Path
from datetime import datetime, timezone
import httpx
from openpyxl import load_workbook
if os.getenv('FOCUS_ALLOW_SYNTHETIC_PRODUCTION_WRITES')!='explicitly-authorized':
    raise SystemExit('Production writes disabled. Use isolated_release_check for normal verification.')
origin='https://ai-focus-3-0.onrender.com'
secret_path=Path('.render-auth/verification.json')
credentials=json.loads(secret_path.read_text())
client=httpx.Client(base_url=origin,timeout=40)
assert client.get('/api/health').json()['status']=='ok'
assert client.get('/api/salons').status_code==401
r=client.post('/api/auth/login',json={k:credentials[k] for k in ['email','password']});r.raise_for_status()
assert 'Secure' in r.headers['set-cookie'] and 'HttpOnly' in r.headers['set-cookie'] and 'SameSite=strict' in r.headers['set-cookie']
client.headers['X-CSRF-Token']=r.json()['csrf'];user=r.json()['user'];assert user['role']=='manager'
assert any(u['email']=='art.gawlik@gmail.com' and u['active'] and u['role']=='manager' for u in client.get('/api/users').json())
name='[TEST TECHNICZNY] Ai Focus 3.0'
existing=client.get('/api/salons',params={'q':name}).json()['items']
if existing: salon=existing[0]
else:
    r=client.post('/api/salons',json={'name':name,'city':'Głogów','notes':'Dane syntetyczne. Test wdrożenia; nie kontaktować się.','potential':0});r.raise_for_status();salon=r.json()
id=salon['id']
assert client.post('/api/salons',json={'name':name,'city':'Głogów'}).status_code==409
assert client.post('/api/salons',json={},headers={'X-CSRF-Token':''}).status_code==403
assert client.post('/api/salons',json={},headers={'Origin':'https://example.invalid'}).status_code==403
assert client.post('/api/salons/'+id+'/activities',json={'kind':'note','body':'[TEST] Zweryfikowano produkcyjny zapis i historię CRM. Brak wysyłki.'}).status_code==201
fields=['name','city','address','contact','email','phone','instagram','territory','specialization','priority','stage','status','potential','next_contact','notes','source_url','version']
payload={k:salon[k] for k in fields};payload['stage']='offer';payload['status']='inactive'
r=client.put('/api/salons/'+id,json=payload);r.raise_for_status();salon=r.json()
assert client.put('/api/salons/'+id,json=payload).status_code==409
r=client.post('/api/appointments',json={'salon_id':id,'starts_at':'2030-01-20T10:00:00+01:00','minutes':45,'status':'proposed','notes':'[TEST TECHNICZNY] Syntetyczny termin do anulowania.'});r.raise_for_status();appointment=r.json()
assert client.post('/api/appointments',json={'salon_id':id,'starts_at':'2030-01-20T10:00:00+01:00','minutes':45,'status':'proposed'}).status_code==409
assert client.patch('/api/appointments/'+appointment['id'],json={'status':'cancelled'}).status_code==200
r=client.post('/api/agent/jobs',json={'salon_id':id,'evidence':'[TEST] Syntetyczny salon testowy w Głogowie. Brak danych prawdziwego klienta.'});r.raise_for_status();job=r.json()
for _ in range(30):
    job=next(j for j in client.get('/api/agent/jobs').json() if j['id']==job['id'])
    if job['status'] in ['completed','failed']:break
    time.sleep(2)
assert job['status']=='completed',job['status']
result=json.loads(job['result']);assert result['sent'] is False and result['engine']=='rules'
wb=load_workbook(io.BytesIO(client.get('/api/export.xlsx').content));assert any(row[0]==name for row in wb.active.iter_rows(min_row=2,values_only=True))
assert client.get('/api/dashboard').json()['total']>=1
backup=client.post('/api/admin/backup',json={'passphrase':credentials['backup_passphrase']},timeout=150)
backup.raise_for_status();assert backup.content.startswith(b'FOCUS3B1')
file=Path('.render-auth/production-backup.enc');file.write_bytes(backup.content);file.chmod(0o600)
state={'salon_id':id,'salon_name':name,'job_id':job['id'],'verification_user_id':user['id'],'checked_at':datetime.now(timezone.utc).isoformat(),'checks':['https','postgres-readiness','manager-login','owner-account','secure-cookie','auth-required','csrf','origin','crm-write-read','duplicate','optimistic-lock','activity-history','appointments-conflict','appointment-cancel','separate-agent-completed','no-send','excel-export','dashboard','encrypted-production-backup']}
Path('.render-auth/production-check-state.json').write_text(json.dumps(state))
Path('test-results').mkdir(exist_ok=True)
Path('test-results/production-check.json').write_text(json.dumps({'origin':origin,'checks':state['checks'],'status':'PASS'},indent=2))
print(json.dumps({'status':'PASS','checks':len(state['checks']),'agent':'completed, rules, sent=false','backup_bytes':len(backup.content),'synthetic_salons':1}))
